#!/usr/bin/env python3
"""M3-lite step 1: inventory + xlsx ingestion.

1) Manifest of all RAG material files (type/size/provenance folder)
2) Unified SKU records from market xlsx databases -> rag/sku-records.jsonl
   (data_only=True reads Excel-cached computed values like 克价)
"""
import glob
import json
import os
import openpyxl

SRC = "/Users/kjonekong/Documents/个人/公众号/普洱茶rag素材"
OUT = "/Users/kjonekong/Documents/puer-ai/rag"
os.makedirs(OUT, exist_ok=True)

def manifest():
    rows = []
    for path in glob.glob(f"{SRC}/**/*", recursive=True):
        if not os.path.isfile(path) or path.endswith(".DS_Store"):
            continue
        rel = os.path.relpath(path, SRC)
        ext = os.path.splitext(path)[1].lstrip(".").lower() or "none"
        rows.append({"path": rel, "ext": ext, "size": os.path.getsize(path),
                     "group": rel.split("/")[0] if "/" in rel else "_root"})
    with open(f"{OUT}/manifest.jsonl", "w", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")
    by_group = {}
    for r in rows:
        g = by_group.setdefault(r["group"], {"files": 0, "exts": {}})
        g["files"] += 1
        g["exts"][r["ext"]] = g["exts"].get(r["ext"], 0) + 1
    print(f"manifest: {len(rows)} files")
    for g, v in sorted(by_group.items()):
        top = sorted(v["exts"].items(), key=lambda x: -x[1])[:4]
        print(f"  {g}: {v['files']} files {dict(top)}")

# column fuzzy mapping to unified schema
COLMAP = {
    "年份": "year", "品名": "name", "形制": "form", "分类": "category",
    "生熟": "raw_or_ripe", "件行情": "market_price_per_jian", "行情/件": "market_price_per_jian", "行情": "market_price",
    "配货价": "alloc_price", "回收价": "buyback_price", "估值": "estimate",
    "最高": "market_high", "规格": "spec_g", "规格2": "pcs_per_jian",
    "件规格": "pcs_per_jian2", "克价": "price_per_g", "rank": "rank",
    "Lot": "lot", "备注": "note", "当年编号": "seq", "编号": "seq",
    "茶行/茶品": "name", "入手时间": "bought_at",
}

def ingest_xlsx():
    out = open(f"{OUT}/sku-records.jsonl", "w", encoding="utf-8")
    total = 0
    for f in sorted(glob.glob(f"{SRC}/*.xlsx")):
        fname = os.path.basename(f)
        try:
            wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        except Exception as e:
            print(f"!! {fname}: {e}")
            continue
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            header_row = None
            cols = {}
            buf = []
            for i, row in enumerate(ws.iter_rows(values_only=True), 1):
                vals = [str(c).strip() if c is not None else "" for c in row or []]
                if header_row is None:
                    hits = sum(1 for v in vals if v in COLMAP)
                    if hits >= 3:
                        header_row = i
                        for j, v in enumerate(vals):
                            if v in COLMAP and v != "":
                                key = COLMAP[v]
                                if key not in cols.values():
                                    cols[j] = key
                    continue
                rec = {}
                for j, key in cols.items():
                    if j < len(vals):
                        rec[key] = vals[j]
                if rec.get("name") and (rec.get("year") or rec.get("market_price") or rec.get("estimate") or rec.get("market_price_per_jian")):
                    rec.update({"source_file": fname, "sheet": sheet, "row": i})
                    buf.append(rec)
            for rec in buf:
                out.write(json.dumps(rec, ensure_ascii=False) + "\n")
            if buf:
                print(f"{fname} [{sheet}]: {len(buf)} records")
            total += len(buf)
        wb.close()
    out.close()
    print(f"TOTAL sku records: {total}")

if __name__ == "__main__":
    manifest()
    ingest_xlsx()
